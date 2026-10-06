"""src.excel_output workbooks from a synthetic summary (no network)."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from src import config
from src.excel_output import THRESHOLD_SUFFIXES, generate_all_workbooks

MONTHS = ["2026-06", "2026-07", "2026-08"]


def _summary() -> pd.DataFrame:
    rows = []
    for r_i, region in enumerate(config.REGIONS):
        for m_i, ym in enumerate(MONTHS):
            total = pd.Period(ym, freq="M").days_in_month * 96
            row = {"REGIONID": region, "YEAR_MONTH": ym, "total_daylight_intervals": total}
            for t_i, suffix in enumerate(THRESHOLD_SUFFIXES):
                count = max(total // 2 - 300 * t_i - 40 * r_i - 10 * m_i, 0)
                row[f"count_below_{suffix}"] = count
                row[f"pct_below_{suffix}"] = round(count / total * 100, 2)
            rows.append(row)
    return pd.DataFrame(rows)


@pytest.fixture(scope="module")
def outdir(tmp_path_factory):
    out = tmp_path_factory.mktemp("xlsx")
    generate_all_workbooks(_summary(), str(out))
    return out


def _pct_cells(ws):
    return [c for row in ws.iter_rows(min_row=2, min_col=2, max_col=1 + len(THRESHOLD_SUFFIXES))
            for c in row]


# --- L2: % cells showed a bare number (25.30) with no unit -------------------

@pytest.mark.parametrize("book, sheet", [
    ("SA_negative_prices.xlsx", "Percentages"),
    ("SA_negative_prices.xlsx", "Heatmap"),
    ("All_States_negative_prices.xlsx", "TAS"),
])
def test_percentage_cells_display_a_percent_sign(outdir, book, sheet):
    ws = load_workbook(outdir / book)[sheet]
    cells = _pct_cells(ws)
    assert len(cells) == len(MONTHS) * len(THRESHOLD_SUFFIXES)
    # The value stays the percentage itself (25.3, as in summary.csv); only the display gains "%".
    assert {c.number_format for c in cells} == {'0.00"%"'}
    assert all(isinstance(c.value, (int, float)) and 0 <= c.value <= 100 for c in cells)


def test_audit_counts_have_no_percent_format(outdir):
    ws = load_workbook(outdir / "SA_negative_prices.xlsx")["Audit"]
    assert all("%" not in c.number_format for row in ws.iter_rows(min_row=2) for c in row)


# --- L2: the heat colours scaled per column, so a 0.37% cell in the -$80 column
# was painted the same red as 74% in the $0 column -----------------------------

def test_heatmap_is_one_fixed_scale_across_every_threshold_column(outdir):
    ws = load_workbook(outdir / "NSW_negative_prices.xlsx")["Heatmap"]
    last_col = chr(ord("A") + len(THRESHOLD_SUFFIXES))
    ranges = [str(cf.sqref) for cf in ws.conditional_formatting]
    assert ranges == [f"B2:{last_col}{len(MONTHS) + 1}"]

    rules = [rule for cf in ws.conditional_formatting for rule in cf.rules]
    assert len(rules) == 1 and rules[0].type == "colorScale"
    cfvo = rules[0].colorScale.cfvo
    # Fixed percentages, not min/percentile/max of the data: a colour means the
    # same share of intervals in every column, every region and every month.
    assert [(v.type, float(v.val)) for v in cfvo] == [("num", 0.0), ("num", 10.0), ("num", 50.0)]
